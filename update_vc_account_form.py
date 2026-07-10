from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SOURCE = Path("/Users/apple/Downloads/个人开户申请表（VC版本）")
OUTPUT = Path("/Users/apple/Documents/VirtuCapital/個人開戶申請表（VC版本-已更新）.docx")


def iter_paragraphs(parent):
    for p in getattr(parent, "paragraphs", []):
        yield p
    for table in getattr(parent, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


def replace_spanning_runs(paragraph, old, new):
    """Replace text even when Word split it across several formatted runs."""
    nodes = list(paragraph._p.iter(qn("w:t")))
    if not nodes:
        return
    while True:
        values = [(n.text or "") for n in nodes]
        combined = "".join(values)
        start = combined.find(old)
        if start < 0:
            return
        end = start + len(old)
        offsets = []
        cursor = 0
        for value in values:
            offsets.append((cursor, cursor + len(value)))
            cursor += len(value)
        first = next(i for i, (a, b) in enumerate(offsets) if a <= start < b)
        last = next(i for i, (a, b) in enumerate(offsets) if a < end <= b)
        first_a, _ = offsets[first]
        last_a, _ = offsets[last]
        prefix = values[first][: start - first_a]
        suffix = values[last][end - last_a :]
        nodes[first].text = prefix + new + suffix
        for i in range(first + 1, last + 1):
            nodes[i].text = ""


def replace_paragraph_text(paragraph, text):
    runs = paragraph.runs
    template_rpr = deepcopy(runs[0]._r.rPr) if runs and runs[0]._r.rPr is not None else None
    for run in list(runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    if template_rpr is not None:
        run._r.insert(0, template_rpr)


doc = Document(SOURCE)

# Replace the institution name throughout the form while retaining surrounding formatting.
parents = [doc]
for section in doc.sections:
    parents.extend([section.header, section.footer, section.first_page_header, section.first_page_footer])
for parent in parents:
    for paragraph in iter_paragraphs(parent):
        replace_spanning_runs(paragraph, "Yellow River Securities Limited", "Virtu Capital Finance Limited")
        replace_spanning_runs(paragraph, "黃河證券有限公司", "Virtu Capital Finance Limited")

# First-page regulatory and contact block. Keep the original short bilingual form.
replace_paragraph_text(
    doc.paragraphs[1],
    "Virtu Capital Finance Limited is registered on the New Zealand Financial Service Providers Register (FSP No.: FSP1011969).",
)
replace_paragraph_text(
    doc.paragraphs[2],
    "Virtu Capital Finance Limited 已在新西蘭金融服務提供商註冊處登記（FSP 編號：FSP1011969）。",
)
replace_paragraph_text(
    doc.paragraphs[3],
    "CARE OF LANE NEAVE, LEVEL 8, 48 SHORTLAND STREET, AUCKLAND CENTRAL, AUCKLAND 1010, NEW ZEALAND",
)
replace_paragraph_text(
    doc.paragraphs[4],
    "新西蘭奧克蘭市中心 SHORTLAND STREET 48 號 8 樓，郵編 1010（轉交 LANE NEAVE）",
)
replace_paragraph_text(doc.paragraphs[5], "Website: www.virtucapital.com")

# Company / licence row in the account-type block.
company_cell = doc.tables[0].rows[2].cells[0]
replace_paragraph_text(company_cell.paragraphs[0], "Virtu Capital Finance Limited (FSP No.: FSP1011969)")
replace_paragraph_text(company_cell.paragraphs[1], "Virtu Capital Finance Limited（FSP 編號：FSP1011969）")

doc.save(OUTPUT)
print(OUTPUT)
